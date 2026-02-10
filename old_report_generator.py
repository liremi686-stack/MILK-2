import os
import openpyxl
from openpyxl.styles import PatternFill, Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter
import shutil
import subprocess
import sys
from datetime import datetime
from collections import defaultdict

class OldReportGenerator:
    def __init__(self, db, date):
        self.db = db
        self.date = date
        self._thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        self._center_align = Alignment(horizontal='center', vertical='center')
        self._month_names = [
            "", "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
            "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"
        ]
        
    def create_report(self):
        """Создает отчет в старом формате (как в govno.py)"""
        try:
            cursor = self.db.conn.cursor()
            
            # Нормализуем формат даты
            if '/' in self.date:
                month, year = self.date.split('/')
            elif '.' in self.date:
                month, year = self.date.split('.')
            else:
                raise ValueError("Неверный формат даты")
            
            month_year = f"{month}/{year}"
            month_year_str = month_year.replace('/', '_')
            
            # Получаем путь к директории программы
            if getattr(sys, 'frozen', False):
                # Если программа запущена как EXE
                program_dir = os.path.dirname(sys.executable)
            else:
                # Если программа запущена как скрипт Python
                program_dir = os.path.dirname(os.path.abspath(__file__))
            
            # Создаем отчет в папке с программой
            main_filename = f"Отчет_за_{month_year_str}.xlsx"
            main_filepath = os.path.join(program_dir, main_filename)
            
            # Создаем книгу Excel
            wb = openpyxl.Workbook()
            
            # Удаляем лист по умолчанию и создаем один лист
            wb.remove(wb.active)
            ws = wb.create_sheet("Отчет")
            
            # Настраиваем стили
            header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
            bold_font = Font(bold=True)
            
            # Заполняем лист
            self._fill_worksheet(ws, month_year, cursor, header_fill, bold_font)
            
            # Сохраняем файл (перезаписываем если существует)
            wb.save(main_filepath)
            
            # Открываем файл
            self._open_excel_file(main_filepath)
            
            return True, f'Отчет создан:\n{main_filename}'
            
        except Exception as e:
            return False, f'Ошибка при создании отчета: {str(e)}'
    
    def _load_all_data(self, cursor, month_year):
        """Загружает все необходимые данные из базы данных"""
        # Загружаем все отделы
        departments = cursor.execute(
            'SELECT id, number, comment FROM departments ORDER BY number'
        ).fetchall()
        
        dept_ids = [dept[0] for dept in departments]
        
        # Единый запрос для ежедневных данных
        placeholders = ','.join('?' * len(dept_ids))
        daily_query = f'''
            SELECT department_id, day, value, color 
            FROM daily_data 
            WHERE department_id IN ({placeholders}) AND date LIKE ?
        '''
        
        daily_data = cursor.execute(daily_query, dept_ids + [f'%/{month_year}']).fetchall()
        
        # Единый запрос для данных рапортов
        reports_query = f'''
            SELECT department_id, 
                   day1, day2, day3, day4, day5,
                   day6, day7, day8, day9, day10,
                   day11, day12, day13, day14, day15,
                   day16, day17, day18, day19, day20,
                   day21, day22, day23, day24, day25,
                   day26, day27, day28, day29, day30, day31
            FROM reports 
            WHERE department_id IN ({placeholders}) AND date LIKE ?
        '''
        
        reports_data = cursor.execute(reports_query, dept_ids + [f'%/{month_year}']).fetchall()
        
        # Единый запрос для месячных данных
        monthly_query = f'''
            SELECT department_id, 
                   decade1_color, decade2_color, decade3_color, total_color
            FROM monthly_data 
            WHERE department_id IN ({placeholders}) AND month_year = ?
        '''
        
        monthly_colors = cursor.execute(monthly_query, dept_ids + [month_year]).fetchall()
        
        # Группируем данные для быстрого доступа
        daily_by_dept = defaultdict(dict)
        for dept_id, day, value, color in daily_data:
            daily_by_dept[dept_id][day] = (value, color)
        
        reports_by_dept = {}
        for row in reports_data:
            dept_id = row[0]
            reports_by_dept[dept_id] = row[1:]  # Пропускаем department_id
        
        colors_by_dept = {}
        for row in monthly_colors:
            dept_id = row[0]
            colors_by_dept[dept_id] = row[1:]
        
        return departments, daily_by_dept, reports_by_dept, colors_by_dept
    
    def _apply_report_changes(self, day_data_dict, reports_data):
        """Применяет изменения из рапортов к данным"""
        if not reports_data:
            return day_data_dict
        
        for day in range(1, 32):
            delta_str = reports_data[day-1]
            if delta_str:
                try:
                    delta = float(delta_str)
                    if delta != 0:
                        current_value, color = day_data_dict.get(day, (0, None))
                        new_value = current_value + delta
                        new_value = max(0, new_value)
                        day_data_dict[day] = (new_value, color)
                except (ValueError, TypeError):
                    pass
        
        return day_data_dict
    
    def _fill_worksheet(self, ws, month_year, cursor, header_fill, bold_font):
        """Оптимизированное заполнение листа Excel"""
        month_num, year = map(int, month_year.split('/'))
        month_name = self._month_names[month_num] if 1 <= month_num <= 12 else ""

        # Первая строка
        ws.merge_cells('A1:AJ1')
        ws['A1'] = f"Сведения за {month_year}"
        ws['A1'].alignment = self._center_align
        ws['A1'].font = Font(bold=True, size=12)

        # Вторая строка
        ws.merge_cells('A2:AJ2')
        ws['A2'] = f'о планируемом** наличии сотрудников ФГКУ "1586" и его структурных подразделений на рабочем месте в {month_name} {year} г. для производства расчёта выдачи молока за вредные условия труда в бухгалтерии ООО "РБЕ"'
        ws['A2'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws['A2'].font = Font(bold=True, size=10)
        ws.row_dimensions[2].height = 42.5  # 1.5 см

        # Третья строка: заголовки
        headers = [
            'Наименование\nподразделения',
            '1', '2', '3', '4', '5', '6', '7', '8', '9', '10', 'итого\n1-10',
            '11', '12', '13', '14', '15', '16', '17', '18', '19', '20', 'итого\n11-20',
            '21', '22', '23', '24', '25', '26', '27', '28', '29', '30', '31', 'итого\n21-31',
            'ИТОГО'
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.fill = header_fill
            cell.border = self._thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.font = Font(bold=True, size=10)
            # Настройка ширины столбцов
            self._set_column_width(ws, col, header)
        
        # Высота строки заголовков
        ws.row_dimensions[3].height = 28.35  # 1 см

        # Загружаем все данные
        departments, daily_by_dept, reports_by_dept, colors_by_dept = self._load_all_data(
            cursor, month_year
        )
        
        # Собираем все данные для расчета итогов
        all_data = []
        row = 4
        
        for dept_id, dept_num, comment in departments:
            # Получаем данные для текущего отдела
            day_data_dict = dict(daily_by_dept.get(dept_id, {}))
            
            # Применяем изменения из рапортов
            reports_data = reports_by_dept.get(dept_id)
            if reports_data:
                day_data_dict = self._apply_report_changes(day_data_dict, reports_data)
            
            # Вычисляем суммы с учетом рапортов
            decade1 = sum(day_data_dict.get(day, (0, None))[0] for day in range(1, 11))
            decade2 = sum(day_data_dict.get(day, (0, None))[0] for day in range(11, 21))
            decade3 = sum(day_data_dict.get(day, (0, None))[0] for day in range(21, 32))
            total = decade1 + decade2 + decade3
            
            # Получаем цвета
            colors = colors_by_dept.get(dept_id, (None, None, None, None))
            
            # Сохраняем данные для итогов
            dept_info = {
                'comment': comment,
                'day_data': day_data_dict,
                'decade1': decade1,
                'decade2': decade2,
                'decade3': decade3,
                'total': total,
                'colors': colors
            }
            all_data.append(dept_info)
            
            # Записываем строку данных
            self._write_department_row(ws, row, comment, day_data_dict, 
                                      decade1, decade2, decade3, total, colors)
            
            row += 1

        # Добавляем строку "ИТОГО"
        total_row = row
        self._write_total_row(ws, total_row, all_data)
        
        # Добавляем примечания
        self._add_notes(ws, total_row)
        
        # Настройки для печати на 2 страницах
        self._setup_print_settings(ws, total_row + 3)  # +3 строки примечаний

    def _set_column_width(self, ws, col, header):
        """Устанавливает ширину столбца в зависимости от его назначения"""
        col_letter = get_column_letter(col)
        
        if col == 1:  # Столбец с названиями отделений
            ws.column_dimensions[col_letter].width = 14.5
        elif col in [12, 23, 35, 36]:  # Столбцы итогов
            ws.column_dimensions[col_letter].width = 8.0
        elif header.isdigit() or header in ['итого\n1-10', 'итого\n11-20', 'итого\n21-31']:
            # Столбцы с днями
            ws.column_dimensions[col_letter].width = 3.0
    
    def _write_department_row(self, ws, row, comment, day_data_dict, 
                             decade1, decade2, decade3, total, colors):
        """Записывает строку данных для одного отдела"""
        # Название подразделения
        name_cell = ws.cell(row=row, column=1, value=comment)
        name_cell.border = self._thin_border
        name_cell.font = Font(size=10)
        name_cell.alignment = Alignment(horizontal='left', vertical='center')
        
        # Данные за каждый день
        for day in range(1, 32):
            value, color = day_data_dict.get(day, (0, None))
            col_num = self._get_day_column(day)
            
            cell = ws.cell(row=row, column=col_num, value=value)
            cell.border = self._thin_border
            cell.font = Font(size=12)
            cell.alignment = self._center_align
            
            # Применяем цвет фона если есть
            self._apply_cell_color(cell, color)
        
        # Итого 1-10
        cell = ws.cell(row=row, column=12, value=decade1)
        cell.border = self._thin_border
        cell.font = Font(size=12)
        cell.alignment = self._center_align
        self._apply_cell_color(cell, colors[0] if colors else None)
        
        # Итого 11-20
        cell = ws.cell(row=row, column=23, value=decade2)
        cell.border = self._thin_border
        cell.font = Font(size=12)
        cell.alignment = self._center_align
        self._apply_cell_color(cell, colors[1] if colors else None)
        
        # Итого 21-31
        cell = ws.cell(row=row, column=35, value=decade3)
        cell.border = self._thin_border
        cell.font = Font(size=12)
        cell.alignment = self._center_align
        self._apply_cell_color(cell, colors[2] if colors else None)
        
        # ВСЕГО
        cell = ws.cell(row=row, column=36, value=total)
        cell.border = self._thin_border
        cell.font = Font(size=12)
        cell.alignment = self._center_align
        self._apply_cell_color(cell, colors[3] if colors else None)
        
        # Высота строки данных
        ws.row_dimensions[row].height = 18.0
    
    def _get_day_column(self, day):
        """Возвращает номер столбца для указанного дня"""
        if day <= 10:
            return day + 1
        elif day <= 20:
            return day + 2
        else:
            return day + 3
    
    def _apply_cell_color(self, cell, color):
        """Применяет цвет фона к ячейке если он задан"""
        if color and color != "#000000" and color != "#00000000":
            try:
                rgb_color = color.lstrip('#')
                if len(rgb_color) == 6 or len(rgb_color) == 8:  # RGB или ARGB
                    cell.fill = PatternFill(
                        start_color=rgb_color, 
                        end_color=rgb_color, 
                        fill_type="solid"
                    )
            except:
                pass  # Игнорируем ошибки цвета
    
    def _write_total_row(self, ws, total_row, all_data):
        """Записывает строку итогов"""
        total_cell = ws.cell(row=total_row, column=1, value="ВСЕГО:")
        total_cell.font = Font(bold=True, size=12)
        total_cell.border = self._thin_border
        
        # Итоги по всем подразделениям
        for day in range(1, 32):
            total_value = sum(
                data['day_data'].get(day, (0, None))[0] 
                for data in all_data
            )
            
            col_num = self._get_day_column(day)
            cell = ws.cell(row=total_row, column=col_num, value=total_value)
            cell.border = self._thin_border
            cell.font = Font(bold=True, size=12)
            cell.alignment = self._center_align
        
        # Итоги по декадам
        total_decade1 = sum(data['decade1'] for data in all_data)
        total_decade2 = sum(data['decade2'] for data in all_data)
        total_decade3 = sum(data['decade3'] for data in all_data)
        total_all = sum(data['total'] for data in all_data)
        
        for col, value in [(12, total_decade1), (23, total_decade2), 
                          (35, total_decade3), (36, total_all)]:
            cell = ws.cell(row=total_row, column=col, value=value)
            cell.border = self._thin_border
            cell.font = Font(bold=True, size=12)
            cell.alignment = self._center_align
    
    def _add_notes(self, ws, total_row):
        """Добавляет примечания в конце таблицы"""
        notes_row = total_row + 1
        
        # Примечание 1
        ws.merge_cells(f'A{notes_row}:AJ{notes_row}')
        note1_cell = ws.cell(
            row=notes_row, column=1, 
            value='*Срок представления Сведений (экз.№2) в ООО АК "Консул" к 29 числу месяца предшествующему выдаче молока**. Данные представлены на основании рапорта (заявления) соответствующих начальников медицинских отделений и других подразделений.'
        )
        note1_cell.font = Font(size=10)
        note1_cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws.row_dimensions[notes_row].height = 28.35
        
        # Примечание 2
        notes_row += 1
        ws.merge_cells(f'A{notes_row}:AJ{notes_row}')
        note2_cell = ws.cell(
            row=notes_row, column=1, 
            value='***Внеплановое отсутствие сотрудников на рабочем месте (болезнь, внеплановый отпуск и т.д.) осуществляются на основании дополнительного рапорта (заявления).'
        )
        note2_cell.font = Font(size=10)
        note2_cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws.row_dimensions[notes_row].height = 28.35
        
        # Примечание 3
        notes_row += 1
        ws.merge_cells(f'A{notes_row}:AJ{notes_row}')
        note3_cell = ws.cell(
            row=notes_row, column=1, 
            value='Начальник отделения продовольственного снабжения ФГКУ "1586 ВКГ"              Е.Акифьева'
        )
        note3_cell.font = Font(bold=True, size=10)
        note3_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[notes_row].height = 28.35
    
    def _setup_print_settings(self, ws, last_row):
        """Настраивает параметры печати"""
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        
        # Устанавливаем область печати
        ws.print_area = f'A1:AJ{last_row}'
        
        # Настройка полей для двусторонней печати
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.3
        ws.page_margins.bottom = 0.3
        ws.page_margins.header = 0.2
        ws.page_margins.footer = 0.2
        
        # Устанавливаем повторяющиеся строки на каждой странице (заголовки)
        ws.print_title_rows = '1:3'
        
        # Настройка верхнего и нижнего колонтитулов
        ws.oddHeader.center.text = "&[File]"
        ws.oddHeader.center.size = 10
        ws.oddFooter.center.text = "Страница &[Page] из &[Pages]"
        ws.oddFooter.center.size = 10
        
        # Для четных страниц (обратная сторона)
        ws.evenHeader.center.text = "&[File]"
        ws.evenHeader.center.size = 10
        ws.evenFooter.center.text = "Страница &[Page] из &[Pages]"
        ws.evenFooter.center.size = 10
    
    def _open_excel_file(self, filename):
        """Открывает файл Excel в ассоциированном приложении"""
        try:
            if sys.platform.startswith('win'):
                os.startfile(filename)
            elif sys.platform.startswith('darwin'):
                subprocess.call(['open', filename])
            else:
                subprocess.call(['xdg-open', filename])
        except Exception as e:
            print(f"Не удалось открыть файл: {e}")