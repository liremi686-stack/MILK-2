import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Side, Font
from openpyxl.utils import get_column_letter
import os
from datetime import datetime
import sqlite3
from contextlib import closing

class ReportGenerator:
    def __init__(self, db_path, progress_callback=None):
        self.db_path = db_path
        self.conn = sqlite3.connect(db_path)
        self.progress_callback = progress_callback
        self._department_mapping = None
        self._thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        self._center_alignment = Alignment(horizontal='center', vertical='center')
        self._font_8pt = Font(size=8)
    
    def update_progress(self, value):
        """Обновление прогресса"""
        if self.progress_callback:
            self.progress_callback(value)
    
    def _get_department_mapping(self):
        """Ленивая загрузка маппинга отделений"""
        if self._department_mapping is None:
            self._department_mapping = {
                '1 - 1 отделение': '1',
                '2 - 2 отделение': '2', 
                '3 - 3 отделение': '3',
                '4 - 4 отделение': '4',
                '5 - 5 отделение': '5',
                '6 - 6 отделение': '6',
                '7 - 7 отделение': '7',
                '8 - 8 отделение': '8',
                '9 - 9 отделение': '9',
                '10 - 10 отделение': '10',
                '11 - 11 отделение': '11',
                '12 - 12 отделение': '12',
                '13 - 13 отделение': '13',
                '14 - 14 отделение': '14',
                '15 - 15 отделение': '15',
                '16 - 16 отделение': '16',
                '17 - 17 отделение': '17',
                '18 - 18 отделение': '18',
                '19 - 19 отделение': '19',
                '20 - 20 отделение': '20',
                '21 - 21 отделение': '21',
                '22 - 22 отделение': '22',
                '23 - 23 отделение': '23',
                '24 - 24 отделение': '24',
                '25 - 25 отделение': '25',
                '26 - 26 отделение': '26',
                '27 - 27 отделение': '27',
                '28 - 28 отделение': '28',
                '29 - 29 отделение': '29',
                '30 - 30 отделение': '30',
                '31 - 31 отделение': '31',
                '32 - 32 отделение': '32',
                '33 - 33 отделение': '33',
                '34 - 34 отделение': '34',
                '35 - 35 отделение': '35',
                '36 - 36 отделение': '36',
                '37 - 37 отделение': '37',
                '38 - 38 отделение': '38',
                '39 - 39 отделение': '39 Экспресс',
                '40 - 40 отделение': '40',
                '41 - 41 отделение': '41',
                '42 - 42 отделение': '42',
                '43 - 43 отделение': 'Функц.Диагн.43 отд.',
                '44 - 44 отделение': '44',
                '45 - 45 отделение': '45(Рентген)',
                '46 - 46 отделение': '46 ультразвук',
                '64 - ВП п.Ильинское': 'ВП п.Ильинское',
                '65 - Приемное отд.': 'Приемное отд.',
                '66 - 66 отделение': '66',
                '67 - Автоотделение': 'Автоотделение',
                '68 - Тестовое': 'Тестовое'
            }
        return self._department_mapping
    
    def get_data_for_month(self, month_year):
        """Оптимизированное получение данных из базы данных"""
        cursor = self.conn.cursor()
        
        # Единый запрос для всех данных
        query = '''
            SELECT 
                d.number,
                d.comment,
                dd.day,
                COALESCE(dd.value, 0) as base_value,
                r.day1, r.day2, r.day3, r.day4, r.day5,
                r.day6, r.day7, r.day8, r.day9, r.day10,
                r.day11, r.day12, r.day13, r.day14, r.day15,
                r.day16, r.day17, r.day18, r.day19, r.day20,
                r.day21, r.day22, r.day23, r.day24, r.day25,
                r.day26, r.day27, r.day28, r.day29, r.day30, r.day31
            FROM departments d
            LEFT JOIN (
                SELECT department_id, day, value 
                FROM daily_data 
                WHERE date LIKE ?
            ) dd ON d.id = dd.department_id
            LEFT JOIN (
                SELECT department_id, 
                    day1, day2, day3, day4, day5,
                    day6, day7, day8, day9, day10,
                    day11, day12, day13, day14, day15,
                    day16, day17, day18, day19, day20,
                    day21, day22, day23, day24, day25,
                    day26, day27, day28, day29, day30, day31
                FROM reports 
                WHERE date LIKE ?
            ) r ON d.id = r.department_id
            ORDER BY d.number
        '''
        
        cursor.execute(query, (f'%/{month_year}', f'%/{month_year}'))
        data = cursor.fetchall()
        
        # Группируем данные по отделам
        dept_data = {}
        current_dept = None
        current_report_days = None
        
        for row in data:
            dept_num, comment, day, base_value, *report_days = row
            
            dept_key = f"{dept_num} - {comment}"
            
            if dept_key not in dept_data:
                dept_data[dept_key] = [0.0] * 31
                current_report_days = report_days
            
            # Обрабатываем ежедневные данные
            if day and 1 <= day <= 31:
                value = float(base_value) if base_value else 0.0
                
                # Применяем изменения из рапортов
                if current_report_days and current_report_days[day-1]:
                    try:
                        delta = float(current_report_days[day-1])
                        value += delta
                        value = max(0, value)
                    except (ValueError, TypeError):
                        pass
                
                dept_data[dept_key][day-1] = value
        
        # Создаем DataFrame
        days_columns = [str(day) for day in range(1, 32)]
        df = pd.DataFrame.from_dict(dept_data, orient='index', columns=days_columns)
        
        return df
    
    def _create_decade_sheet(self, wb, sheet_name, days_range, month, year):
        """Создает лист для декады"""
        sheet = wb.create_sheet(title=sheet_name)
        
        # Настройка страницы
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.orientation = sheet.ORIENTATION_PORTRAIT
        
        # Установка полей (в дюймах)
        margins = sheet.page_margins
        margins.left = 0.5 / 2.54
        margins.right = 0.5 / 2.54
        margins.top = 1.0 / 2.54
        margins.bottom = 1.0 / 2.54
        margins.header = 0.5 / 2.54
        margins.footer = 0.5 / 2.54
        
        sheet.sheet_properties.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        
        return sheet
    
    def _setup_decade_headers(self, sheet, days_range, month, year):
        """Настраивает заголовки для декады"""
        start_day = days_range[0]
        end_day = days_range[-1]
        
        # Первая строка заголовка
        last_col_letter = 'N' if len(days_range) == 11 else 'M'
        merge_range = f'A1:{last_col_letter}1'
        sheet.merge_cells(merge_range)
        sheet['A1'] = f'СВОДНАЯ ВЕДОМОСТЬ C {start_day}-{end_day} {month}.{year} г.'
        sheet['A1'].alignment = self._center_alignment
        sheet['A1'].font = Font(bold=True, size=10)
        
        # Вторая строка заголовка
        merge_range = f'A2:{last_col_letter}2'
        sheet.merge_cells(merge_range)
        sheet['A2'] = 'на выдачу молока сотрудникам 1586 ВКГ занятым во вредных условиях труда'
        sheet['A2'].alignment = self._center_alignment
        sheet['A2'].font = Font(bold=True, size=10)
        
        # Заголовки столбцов
        headers = ['Отделение'] + [str(day) for day in days_range] + ['Всего', 'фамилия получателя', 'Подпись']
        
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=3, column=col, value=header)
            cell.alignment = self._center_alignment
            cell.font = self._font_8pt
            cell.border = self._thin_border
        
        return len(headers)
    
    def _fill_decade_data(self, sheet, df, days_range, start_row, total_columns):
        """Заполняет данные на листе декады"""
        row_idx = start_row
        department_mapping = self._get_department_mapping()
        
        for full_dept_name in df.index:
            short_name = department_mapping.get(full_dept_name, full_dept_name)
            dept_data = df.loc[full_dept_name]
            
            # Название отделения
            cell = sheet.cell(row=row_idx, column=1, value=short_name)
            cell.alignment = self._center_alignment
            cell.font = self._font_8pt
            cell.border = self._thin_border
            
            # Данные по дням
            for col_idx, day in enumerate(days_range, 2):
                day_str = str(day)
                if day_str in dept_data.index:
                    value = dept_data[day_str]
                    if pd.isna(value) or value == 0:
                        cell_value = ''
                    else:
                        cell_value = value / 2  # Делим на 2 как в оригинале
                else:
                    cell_value = ''
                
                cell = sheet.cell(row=row_idx, column=col_idx, value=cell_value)
                cell.alignment = self._center_alignment
                cell.font = self._font_8pt
                cell.border = self._thin_border
            
            # Столбец "Всего" с формулой
            total_col = len(days_range) + 2
            start_col_letter = get_column_letter(2)
            end_col_letter = get_column_letter(len(days_range) + 1)
            cell = sheet.cell(
                row=row_idx, column=total_col,
                value=f'=SUM({start_col_letter}{row_idx}:{end_col_letter}{row_idx})'
            )
            cell.alignment = self._center_alignment
            cell.font = self._font_8pt
            cell.border = self._thin_border
            
            # Пустые ячейки для фамилии и подписи
            for i in range(2):
                cell = sheet.cell(row=row_idx, column=total_col + i + 1, value='')
                cell.alignment = self._center_alignment
                cell.font = self._font_8pt
                cell.border = self._thin_border
            
            # Высота строки
            sheet.row_dimensions[row_idx].height = 20
            row_idx += 1
        
        return row_idx
    
    def _add_decade_totals(self, sheet, days_range, total_row, total_columns):
        """Добавляет строку итогов для декады"""
        # Заголовок "ВСЕГО:"
        cell = sheet.cell(row=total_row, column=1, value='ВСЕГО:')
        cell.alignment = self._center_alignment
        cell.font = self._font_8pt
        cell.border = self._thin_border
        
        # Итоги по дням
        for col_idx, day in enumerate(days_range, 2):
            col_letter = get_column_letter(col_idx)
            cell = sheet.cell(
                row=total_row, column=col_idx,
                value=f'=SUM({col_letter}4:{col_letter}{total_row-1})'
            )
            cell.alignment = self._center_alignment
            cell.font = self._font_8pt
            cell.border = self._thin_border
        
        # Итог по столбцу "Всего"
        total_col = len(days_range) + 2
        total_col_letter = get_column_letter(total_col)
        cell = sheet.cell(
            row=total_row, column=total_col,
            value=f'=SUM({total_col_letter}4:{total_col_letter}{total_row-1})'
        )
        cell.alignment = self._center_alignment
        cell.font = self._font_8pt
        cell.border = self._thin_border
        
        # Пустые ячейки для фамилии и подписи
        for i in range(2):
            cell = sheet.cell(row=total_row, column=total_col + i + 1, value='')
            cell.alignment = self._center_alignment
            cell.font = self._font_8pt
            cell.border = self._thin_border
        
        sheet.row_dimensions[total_row].height = 20
    
    def _setup_decade_columns(self, sheet, days_range, total_columns):
        """Настраивает ширину столбцов для декады"""
        sheet.column_dimensions['A'].width = 22.0
        
        # Столбцы дней
        for col in range(2, len(days_range) + 2):
            col_letter = get_column_letter(col)
            sheet.column_dimensions[col_letter].width = 4.0
        
        # Столбец "Всего"
        total_col = len(days_range) + 2
        total_col_letter = get_column_letter(total_col)
        sheet.column_dimensions[total_col_letter].width = 6.0
        
        # Столбцы с фамилией и подписью
        for i in range(2):
            col = total_col + i + 1
            col_letter = get_column_letter(col)
            width = 18.0 if i == 0 else 12.0
            sheet.column_dimensions[col_letter].width = width
    
    def _add_signatures(self, sheet, total_row):
        """Добавляет подписи в конце таблицы"""
        signature_row = total_row + 2
        
        # Зам. Управляющего столовой
        sheet.merge_cells(f'A{signature_row}:C{signature_row}')
        cell = sheet.cell(
            row=signature_row, column=1,
            value='Зам. Управляющего столовой ООО Р Б Е'
        )
        cell.alignment = self._center_alignment
        cell.font = self._font_8pt
        
        # Т. Сурова
        sheet.merge_cells(f'D{signature_row}:F{signature_row}')
        cell = sheet.cell(row=signature_row, column=4, value='Т. Сурова')
        cell.alignment = self._center_alignment
        cell.font = self._font_8pt
        
        # Выдал
        cell = sheet.cell(row=signature_row + 1, column=1, value='Выдал:')
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.font = self._font_8pt
        
        # Бухгалтер
        cell = sheet.cell(row=signature_row + 2, column=1, value='Бухгалтер:')
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.font = self._font_8pt
        
        # Высота строк с подписями
        for row in [signature_row, signature_row + 1, signature_row + 2]:
            sheet.row_dimensions[row].height = 20
    
    def create_svodka(self, month_year, output_filename=None):
        """Создает файл сводки из данных базы данных"""
        try:
            self.update_progress(0)
            
            # Приведение month_year к единому формату
            if '/' in month_year:
                month, year = month_year.split('/')
            elif '.' in month_year:
                month, year = month_year.split('.')
            else:
                raise ValueError("Неверный формат month_year. Ожидается MM/YYYY или MM.YYYY")
            
            if output_filename is None:
                output_filename = f'Сводка_{month}_{year}.xlsx'
            
            # Проверяем существование файла
            if os.path.exists(output_filename):
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                output_filename = f'Сводка_{month}_{year}_{timestamp}.xlsx'
            
            print(f"\nОбработка данных за {month}.{year}")
            print(f"Создание сводки: {output_filename}")
            
            self.update_progress(10)
            
            # Получаем данные из базы данных
            df_data = self.get_data_for_month(f"{month}/{year}")
            
            self.update_progress(30)
            
            # Создаем Excel-файл
            wb = openpyxl.Workbook()
            wb.remove(wb.active)  # Удаляем лист по умолчанию
            
            # Создаем листы для трех декад
            decades = [
                ('1 декада', range(1, 11), 'M'),
                ('2 декада', range(11, 21), 'M'),
                ('3 декада', range(21, 32), 'N')
            ]
            
            for sheet_name, days_range, last_col_letter in decades:
                self.update_progress(40 + (decades.index((sheet_name, days_range, last_col_letter)) * 10))
                
                sheet = self._create_decade_sheet(wb, sheet_name, days_range, month, year)
                total_columns = self._setup_decade_headers(sheet, days_range, month, year)
                
                # Заполняем данные
                last_row = self._fill_decade_data(sheet, df_data, days_range, 4, total_columns)
                
                # Добавляем итоги
                self._add_decade_totals(sheet, days_range, last_row, total_columns)
                
                # Настраиваем столбцы
                self._setup_decade_columns(sheet, days_range, total_columns)
                
                # Добавляем подписи
                self._add_signatures(sheet, last_row)
            
            self.update_progress(80)
            
            # Сохраняем файл
            wb.save(output_filename)
            
            self.update_progress(100)
            
            print(f"  ✓ Создан файл: {output_filename}")
            return output_filename
            
        except Exception as e:
            print(f"  ✗ Ошибка при создании сводки: {str(e)}")
            import traceback
            traceback.print_exc()
            return None
    
    def close(self):
        """Закрывает соединение с базой данных"""
        if hasattr(self, 'conn') and self.conn:
            self.conn.close()

# Функция для быстрого создания сводки
def create_svodka_from_database(db_path, month_year):
    """Создает сводку из данных базы данных за указанный месяц"""
    with closing(ReportGenerator(db_path)) as generator:
        return generator.create_svodka(month_year)

# Основная программа для обработки
def main():
    """
    Основная функция для обработки данных из базы данных
    """
    import sys
    from config import DB_PATH
    
    print("=" * 60)
    print("СОЗДАНИЕ СВОДОК ИЗ БАЗЫ ДАННЫХ MILKMAN")
    print("=" * 60)
    
    if len(sys.argv) > 1:
        month_year = sys.argv[1]
    else:
        month_year = input("Введите месяц и год в формате MM/YYYY или MM.YYYY: ")
    
    if not month_year:
        print("Месяц и год не указаны!")
        return
    
    result = create_svodka_from_database(DB_PATH, month_year)
    
    if result:
        print(f"\nСводка успешно создана: {result}")
    else:
        print("\nНе удалось создать сводку")

if __name__ == "__main__":
    main()